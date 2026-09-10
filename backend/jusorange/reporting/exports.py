"""Exports Excel des rapports.

Chaque classeur s'ouvre sur une feuille « Synthèse » : les chiffres clés avec
leur évolution, ce qu'il faut en retenir, et deux graphiques. Les feuilles de
détail qui suivent restent les tableaux bruts, pour qui veut vérifier ligne à
ligne. Un tableau seul oblige le lecteur à refaire l'analyse dans sa tête —
et hors de l'application, personne ne la refait.
"""
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from datetime import date

# Styles réutilisables
THIN_BORDER = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000'),
)
ORANGE_FILL = PatternFill(start_color='FF6A3A', end_color='FF6A3A', fill_type='solid')
ORANGE_HEADER_FONT = Font(bold=True, color='FFFFFF', size=11)

# Fonds des constats. La couleur double toujours un libellé de gravité écrit
# en clair : sur une impression noir et blanc, l'information doit survivre.
FONDS_TON = {
    'positif': PatternFill('solid', start_color='E3F5EC'),
    'attention': PatternFill('solid', start_color='FDF3DC'),
    'alerte': PatternFill('solid', start_color='FBE3E3'),
    'neutre': PatternFill('solid', start_color='F1F1F0'),
}
LIBELLES_TON = {
    'positif': 'Favorable',
    'attention': 'À surveiller',
    'alerte': 'Alerte',
    'neutre': 'Constat',
}
# Mêmes teintes que les graphiques de l'application, pour que le classeur et
# l'écran racontent visiblement la même histoire.
COULEURS_SERIES = ['EB6834', '2A78D6', '1BAF7A', '4A3AA7', 'EDA100']


def _set_column_widths(ws, width=15):
    """Définit la largeur des colonnes sans toucher aux MergedCell."""
    for col_idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _format_table(ws, start_row, end_row, num_cols, header_row_offset=None):
    """Applique bordures à une plage. Si header_row_offset (0=1ère ligne), applique fond orange."""
    for row in range(start_row, end_row + 1):
        for col in range(1, num_cols + 1):
            try:
                cell = ws.cell(row=row, column=col)
                if not hasattr(cell, 'column_letter'):
                    continue
                cell.border = THIN_BORDER
                if header_row_offset is not None and row == start_row + header_row_offset:
                    cell.fill = ORANGE_FILL
                    cell.font = ORANGE_HEADER_FONT
            except (AttributeError, TypeError):
                pass


def _titre_section(ws, texte, colonnes=6):
    """Bandeau de section, sur sa propre ligne."""
    ws.append([])
    ligne = ws.max_row + 1
    ws.cell(row=ligne, column=1, value=texte).font = Font(bold=True, size=12, color='FF6A3A')
    ws.append([])
    return ligne


def _entetes(ws, colonnes):
    """Ligne d'en-tête au style du reste des exports."""
    ws.append(colonnes)
    ligne = ws.max_row
    for i in range(1, len(colonnes) + 1):
        c = ws.cell(row=ligne, column=i)
        c.fill = ORANGE_FILL
        c.font = ORANGE_HEADER_FONT
        c.border = THIN_BORDER
    return ligne


def _variation_texte(comp):
    """« +18,4 % » ou « nouveau » — jamais un pourcentage inventé.

    Partir de zéro n'est pas une hausse de 100 % : c'est une apparition, et
    l'écrire en pourcentage tromperait le lecteur du classeur comme celui de
    l'écran.
    """
    if not comp:
        return ''
    v = comp.get('variation')
    if v is None:
        return 'nouveau' if comp.get('valeur') else '—'
    if abs(v) < 0.05:
        return 'stable'
    return f"{v:+.1f} %"


def ajouter_feuille_synthese(wb, titre, date_debut, date_fin, analyse, kpis=None):
    """Insère en première position la feuille de lecture du rapport.

    `analyse` est le bloc calculé par reporting.services.analytique ; `kpis`
    est la liste (libellé, clé de comparaison) des chiffres de tête propres au
    module. Sans `analyse`, la feuille n'est pas créée : mieux vaut pas de
    synthèse qu'une synthèse vide.
    """
    if not analyse:
        return None

    ws = wb.create_sheet('Synthèse', 0)

    ws['A1'] = f"{titre} — {date_debut.strftime('%d/%m/%Y')} → {date_fin.strftime('%d/%m/%Y')}"
    ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    ws['A1'].fill = ORANGE_FILL
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:H1')

    prec = analyse.get('periode_precedente') or {}
    if prec:
        ws['A2'] = (f"Comparé à la période précédente de même durée : "
                    f"{prec['date_debut']} → {prec['date_fin']}")
        ws['A2'].font = Font(italic=True, size=10, color='666666')

    # --- Ce qu'il faut retenir : l'interprétation, en toutes lettres ---
    faits = analyse.get('faits') or []
    if faits:
        _titre_section(ws, "Ce qu'il faut retenir")
        _entetes(ws, ['Niveau', 'Constat'])
        for f in faits:
            ws.append([LIBELLES_TON.get(f['ton'], 'Constat'), f['texte']])
            ligne = ws.max_row
            fond = FONDS_TON.get(f['ton'], FONDS_TON['neutre'])
            for col in (1, 2):
                ws.cell(row=ligne, column=col).fill = fond
                ws.cell(row=ligne, column=col).border = THIN_BORDER
            ws.cell(row=ligne, column=2).alignment = Alignment(wrap_text=True, vertical='top')

    # --- Chiffres clés et leur évolution ---
    comparaison = analyse.get('comparaison') or {}
    if kpis and comparaison:
        _titre_section(ws, 'Chiffres clés')
        _entetes(ws, ['Indicateur', 'Période', 'Période précédente', 'Évolution'])
        for libelle, cle in kpis:
            comp = comparaison.get(cle)
            if not comp:
                continue
            ws.append([libelle, comp['valeur'], comp['precedent'], _variation_texte(comp)])
            for col in range(1, 5):
                ws.cell(row=ws.max_row, column=col).border = THIN_BORDER

    # --- Indicateurs dérivés : les ratios que les totaux ne donnent pas ---
    indicateurs = [i for i in (analyse.get('indicateurs') or []) if i.get('valeur') is not None]
    if indicateurs:
        _titre_section(ws, 'Indicateurs dérivés')
        _entetes(ws, ['Indicateur', 'Valeur', 'Unité', 'Lecture'])
        for i in indicateurs:
            ws.append([i['label'], i['valeur'], (i.get('unite') or '').strip(), i.get('aide', '')])
            for col in range(1, 5):
                ws.cell(row=ws.max_row, column=col).border = THIN_BORDER

    # --- Évolution + graphique en courbes ---
    serie = analyse.get('serie') or []
    mesures = analyse.get('serie_mesures') or []
    if serie and mesures:
        _titre_section(ws, 'Évolution sur la période')
        entete = _entetes(ws, ['Période'] + [m['label'] for m in mesures])
        premiere = ws.max_row + 1
        for point in serie:
            ws.append([point['periode']] + [point.get(m['cle'], 0) for m in mesures])
            for col in range(1, len(mesures) + 2):
                ws.cell(row=ws.max_row, column=col).border = THIN_BORDER
        derniere = ws.max_row

        courbe = LineChart()
        courbe.title = 'Évolution sur la période'
        courbe.style = 2
        courbe.height, courbe.width = 8, 18
        # Une seule échelle : deux mesures d'ordres de grandeur différents ne
        # partagent jamais un axe secondaire, la comparaison visuelle y serait
        # fausse.
        donnees = Reference(ws, min_col=2, max_col=1 + len(mesures),
                            min_row=entete, max_row=derniere)
        libelles = Reference(ws, min_col=1, min_row=premiere, max_row=derniere)
        courbe.add_data(donnees, titles_from_data=True)
        courbe.set_categories(libelles)
        # openpyxl marque les axes « supprimés » par défaut : Excel les masque,
        # et le graphique perd ses dates et ses valeurs. Sans cela, la courbe
        # est un trait sans repère.
        courbe.x_axis.delete = False
        courbe.y_axis.delete = False
        courbe.y_axis.majorGridlines = None
        for i, serie_graph in enumerate(courbe.series):
            serie_graph.graphicalProperties.line.solidFill = COULEURS_SERIES[i % len(COULEURS_SERIES)]
            serie_graph.graphicalProperties.line.width = 22000  # ~2 pt
            serie_graph.smooth = False
        # Une légende n'a de sens qu'à partir de deux séries : le titre nomme
        # déjà la mesure quand elle est seule.
        if len(mesures) < 2:
            courbe.legend = None
        ws.add_chart(courbe, f'F{entete}')

    # --- Répartition + graphique en barres ---
    conc = analyse.get('concentration') or {}
    lignes_conc = conc.get('lignes') or []
    if lignes_conc:
        titre_conc = analyse.get('concentration_titre', 'Répartition')
        ligne_titre = _titre_section(ws, titre_conc)
        if conc.get('nb_acteurs', 0) > 1:
            # Sur la ligne SUIVANTE : écrire sur `max_row` retombait sur le
            # titre lui-même et l'effaçait.
            ws.cell(
                row=ligne_titre + 1, column=1,
                value=(f"{conc['acteurs_80pct']} sur {conc['nb_acteurs']} font 80 % du total"
                       + (f" — les 3 premiers en font {conc['part_top3']:.0f} %"
                          if conc.get('part_top3') is not None else ''))
            ).font = Font(italic=True, size=10, color='666666')
        entete = _entetes(ws, ['Libellé', 'Valeur', 'Part (%)', 'Cumul (%)'])
        premiere = ws.max_row + 1
        for l in lignes_conc:
            ws.append([l['libelle'], l['valeur'], l['part'], l['cumul']])
            for col in range(1, 5):
                ws.cell(row=ws.max_row, column=col).border = THIN_BORDER
        derniere = ws.max_row

        if len(lignes_conc) > 1:
            barres = BarChart()
            barres.type = 'bar'  # horizontales : les libellés sont des noms
            barres.title = titre_conc
            barres.height, barres.width = max(6, len(lignes_conc) * 1.1), 18
            barres.add_data(Reference(ws, min_col=2, min_row=entete, max_row=derniere),
                            titles_from_data=True)
            barres.set_categories(Reference(ws, min_col=1, min_row=premiere, max_row=derniere))
            # Idem : sans axes, le Pareto perd les noms qu'il est censé classer.
            barres.x_axis.delete = False
            barres.y_axis.delete = False
            barres.gapWidth = 40
            # Excel empile les catégories du bas vers le haut : le plus gros
            # contributeur apparaît donc en bas. Inverser l'axe retournait
            # aussi l'échelle des valeurs (700 000 à gauche, 0 à droite) et
            # vidait les barres de leur couleur — on garde le rendu natif,
            # lisible, et le tableau juste au-dessus donne l'ordre exact.
            # Une seule teinte : les barres mesurent la même grandeur, la
            # couleur ne code donc aucune identité.
            barres.series[0].graphicalProperties.solidFill = COULEURS_SERIES[0]
            barres.legend = None
            ws.add_chart(barres, f'F{entete}')

    ws.column_dimensions['A'].width = 34
    for lettre in ('B', 'C', 'D'):
        ws.column_dimensions[lettre].width = 20
    ws.column_dimensions['E'].width = 4
    return ws


def export_recolte_excel(date_debut, date_fin, cueillettes, par_producteur=None, par_zone=None,
                         analyse=None):
    """Exporte le rapport récolte en Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport Récolte"
    ws['A1'] = f"Rapport Récolte — {date_debut.strftime('%d/%m/%Y')} → {date_fin.strftime('%d/%m/%Y')}"
    ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    ws['A1'].fill = ORANGE_FILL
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:G1')
    ws.append([])
    headers = ['Producteur', 'Date', 'Qté totale (kg)', 'Qté bonne (kg)', 'Qté mauvaise (kg)', 'Taux qualité (%)', 'Observation']
    ws.append(headers)
    for c in cueillettes:
        t = round((c.qte_bon / c.qte_total * 100), 2) if c.qte_total else 0
        ws.append([c.get_producteur_display(), c.date_cueil, c.qte_total, c.qte_bon, c.qte_mauvais, t, c.observation or ''])
    # Tableau principal : lignes 3 (en-tête) à 3+len(cueillettes)
    _format_table(ws, 3, 3 + len(cueillettes), 7, header_row_offset=0)
    if par_producteur:
        ws.append([])
        ws.append(['Par producteur'])
        ws.append(['Producteur', 'Nb', 'Qté totale', 'Qté bonne'])
        start_row = ws.max_row
        for p in par_producteur:
            ws.append([p.get('producteur__nom_complet', ''), p.get('nb', 0), p.get('qte_total', 0), p.get('qte_bon', 0)])
        _format_table(ws, start_row, ws.max_row, 4, header_row_offset=0)
    if par_zone:
        ws.append([])
        ws.append(['Par zone'])
        ws.append(['Zone', 'Nb', 'Qté totale', 'Qté bonne'])
        start_row = ws.max_row
        for z in par_zone:
            ws.append([z.get('producteur__zone', ''), z.get('nb', 0), z.get('qte_total', 0), z.get('qte_bon', 0)])
        _format_table(ws, start_row, ws.max_row, 4, header_row_offset=0)
    ajouter_feuille_synthese(
        wb, 'Rapport Récolte', date_debut, date_fin, analyse,
        kpis=[('Cueillettes', 'nb_cueillettes'), ('Récolte totale (kg)', 'qte_total_kg'), ('Dont bonnes (kg)', 'qte_bon_kg'), ('Pertes (kg)', 'qte_mauvais_kg')])
    _set_column_widths(ws, 15)
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="rapport_recolte_{date_debut}_{date_fin}.xlsx"'
    wb.save(response)
    return response


def export_appro_excel(date_debut, date_fin, receptions, evolution_stock=None, stock_actuel=None,
                       analyse=None):
    """Exporte le rapport appro en Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport Réceptions"
    ws['A1'] = f"Rapport Réceptions — {date_debut.strftime('%d/%m/%Y')} → {date_fin.strftime('%d/%m/%Y')}"
    ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    ws['A1'].fill = ORANGE_FILL
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:F1')
    ws.append([])
    if evolution_stock:
        ws.append(['Évolution réceptions (oranges kg/mois)'])
        ws.append(['Mois', 'Qté bonne'])
        start_row = ws.max_row
        for e in evolution_stock:
            ws.append([e['mois'], e['qte_bon']])
        _format_table(ws, start_row, ws.max_row, 2, header_row_offset=0)
        ws.append([])
    if stock_actuel:
        ws.append(['Stock actuel'])
        ws.append(['Article', 'Quantité', 'Seuil'])
        start_row = ws.max_row
        for a in stock_actuel:
            ws.append([a.get_type_art_display(), a.qte_art, a.seuil_alerte])
        _format_table(ws, start_row, ws.max_row, 3, header_row_offset=0)
        ws.append([])
    headers = ['N° réception', 'Date', 'Cueillette', 'Qté reçue (kg)', 'Qté bonne (kg)', 'Lieu dépôt']
    ws.append(headers)
    for r in receptions:
        ws.append([r.num_recp, r.date_recp, str(r.get_cueillette_display())[:50], r.qte_recue, r.qte_bon, r.lieu_depot])
    _format_table(ws, ws.max_row - len(receptions), ws.max_row, 6, header_row_offset=0)
    ajouter_feuille_synthese(
        wb, 'Rapport Approvisionnement', date_debut, date_fin, analyse,
        kpis=[('Réceptions', 'nb_receptions'), ('Quantité reçue (kg)', 'qte_recue_kg'), ('Dont bonnes (kg)', 'qte_bon_kg'), ('Pertes (kg)', 'qte_mauvais_kg')])
    _set_column_widths(ws, 18)
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="rapport_appro_{date_debut}_{date_fin}.xlsx"'
    wb.save(response)
    return response


def export_fabrication_excel(date_debut, date_fin, productions, analyse=None):
    """Exporte le rapport fabrication en Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport Production"
    ws['A1'] = f"Rapport Production — {date_debut.strftime('%d/%m/%Y')} → {date_fin.strftime('%d/%m/%Y')}"
    ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    ws['A1'].fill = ORANGE_FILL
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:F1')
    ws.append([])
    headers = ['N° OF', 'Date', 'Recette', 'Statut', 'Volume (L)', 'Test qualité']
    ws.append(headers)
    for p in productions:
        ws.append([p.numero_of, p.date_of, p.get_recette_display(), p.get_statut_production_display(), p.volume_final_l, p.get_test_qualite_display() or ''])
    _format_table(ws, 3, 3 + len(productions), 6, header_row_offset=0)
    ajouter_feuille_synthese(
        wb, 'Rapport Fabrication', date_debut, date_fin, analyse,
        kpis=[('Productions', 'nb_productions'), ('Volume total (L)', 'volume_total_l'), ('Terminées', 'nb_terminees')])
    _set_column_widths(ws, 15)
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="rapport_fabrication_{date_debut}_{date_fin}.xlsx"'
    wb.save(response)
    return response


def export_emballage_excel(date_debut, date_fin, conditionnements, bouteilles_par_statut=None,
                           bouteilles=None, analyse=None):
    """Exporte le rapport emballage en Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport Conditionnement"
    ws['A1'] = f"Rapport Conditionnement — {date_debut.strftime('%d/%m/%Y')} → {date_fin.strftime('%d/%m/%Y')}"
    ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    ws['A1'].fill = ORANGE_FILL
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:G1')
    ws.append([])
    if bouteilles_par_statut:
        ws.append(['Bouteilles par statut'])
        ws.append(['Statut', 'Quantité'])
        start_row = ws.max_row
        for b in bouteilles_par_statut:
            ws.append([b['label'], b['count']])
        _format_table(ws, start_row, ws.max_row, 2, header_row_offset=0)
        ws.append([])
    if bouteilles:
        ws.append(['Bouteilles (échantillon)'])
        ws.append(['ID', 'Conditionnement', 'Format', 'Statut', 'DLC'])
        start_row = ws.max_row
        for b in bouteilles:
            ws.append([b.pk, b.conditionnement.numero_cond, b.get_format_display(), b.get_statut_stock_display(), b.dlc])
        _format_table(ws, start_row, ws.max_row, 5, header_row_offset=0)
        ws.append([])
    headers = ['N° conditionnement', 'Date', 'Production', '33cl', '1L', 'Volume (L)', 'DLC']
    ws.append(headers)
    for c in conditionnements:
        ws.append([c.numero_cond, c.date_cond, c.production.numero_of, c.qte_33cl, c.qte_1l, c.volume_utilisee, c.dlc])
    _format_table(ws, ws.max_row - len(conditionnements), ws.max_row, 7, header_row_offset=0)
    ajouter_feuille_synthese(
        wb, 'Rapport Emballage', date_debut, date_fin, analyse,
        kpis=[('Conditionnements', 'nb_conditionnements'), ('Bouteilles 33cl', 'qte_33cl'), ('Bouteilles 1L', 'qte_1l'), ('Volume utilisé (L)', 'volume_l')])
    _set_column_widths(ws, 15)
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="rapport_emballage_{date_debut}_{date_fin}.xlsx"'
    wb.save(response)
    return response


def export_entrepot_excel(date_debut, date_fin, inventaires, analyse=None):
    """Exporte le rapport entrepot en Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport Inventaire"
    ws['A1'] = f"Rapport Inventaire — {date_debut.strftime('%d/%m/%Y')} → {date_fin.strftime('%d/%m/%Y')}"
    ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    ws['A1'].fill = ORANGE_FILL
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:G1')
    ws.append([])
    headers = ['Date', 'Article', 'Qté système', 'Qté dépôt', 'Écart', 'Statut', 'Qualité']
    ws.append(headers)
    for i in inventaires:
        ws.append([i.date_inv, i.article.get_type_art_display(), i.qte_systeme, i.qte_depot, i.ecart, i.get_statut_display(), i.get_qualite_display()])
    _format_table(ws, 3, 3 + len(inventaires), 7, header_row_offset=0)
    ajouter_feuille_synthese(
        wb, 'Rapport Entrepôt', date_debut, date_fin, analyse,
        kpis=[('Inventaires', 'nb_inventaires'), ('Bloqués', 'nb_bloque'), ('Qualité mauvaise', 'nb_qualite_mauvais')])
    _set_column_widths(ws, 15)
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="rapport_entrepot_{date_debut}_{date_fin}.xlsx"'
    wb.save(response)
    return response


def export_distribution_excel(date_debut, date_fin, ventes, commandes, factures=None, paiements=None,
                              receptions_tresorerie=None, analyse=None):
    """Exporte le rapport distribution en Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport Commercialisation"
    ws['A1'] = f"Rapport Commercialisation — {date_debut.strftime('%d/%m/%Y')} → {date_fin.strftime('%d/%m/%Y')}"
    ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    ws['A1'].fill = ORANGE_FILL
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:G1')
    ws.append([])
    ws.append(['Ventes'])
    ws.append(['Date', 'Client', 'Montant (XOF)', 'Statut'])
    start_row = ws.max_row
    for v in ventes:
        ws.append([v.date_vente, v.client.nom_complet, v.montant_total, v.get_statut_paiement_display()])
    _format_table(ws, start_row, ws.max_row, 4, header_row_offset=0)
    ws.append([])
    ws.append(['Commandes'])
    ws.append(['Date', 'Client', '33cl', '1L', 'Vente'])
    start_row = ws.max_row
    for c in commandes:
        v = f"#{c.vente_id}" if c.vente_id else "En attente"
        ws.append([c.date_cmd, c.client.nom_complet, c.quantite_33cl, c.quantite_1l, v])
    _format_table(ws, start_row, ws.max_row, 5, header_row_offset=0)
    if factures:
        ws.append([])
        ws.append(['Factures'])
        ws.append(['N°', 'Date', 'Client', 'Montant', 'Statut'])
        start_row = ws.max_row
        for f in factures:
            ws.append([f.num_fact, f.date_fact, f.vente.client.nom_complet, f.montant, f.get_statut_display()])
        _format_table(ws, start_row, ws.max_row, 5, header_row_offset=0)
    if paiements:
        ws.append([])
        ws.append(['Paiements'])
        ws.append(['Date', 'Client', 'Montant', 'Mode'])
        start_row = ws.max_row
        for p in paiements:
            ws.append([p.date_paie, p.facture.vente.client.nom_complet, p.montant, p.get_mode_paie_display()])
        _format_table(ws, start_row, ws.max_row, 4, header_row_offset=0)
    if receptions_tresorerie:
        ws.append([])
        ws.append(['Trésorerie'])
        ws.append(['Date', 'Client', 'Déclaré', 'Reçu', 'Écart', 'Traité', 'Observation / Justification'])
        start_row = ws.max_row
        for r in receptions_tresorerie:
            ws.append([
                r.date_reception,
                r.paiement.facture.vente.client.nom_complet,
                r.paiement.montant,
                r.montant_recu,
                r.ecart,
                'Oui' if r.ecart_traite else 'Non',
                r.observation or '',
            ])
        _format_table(ws, start_row, ws.max_row, 7, header_row_offset=0)
    ajouter_feuille_synthese(
        wb, 'Rapport Commercialisation', date_debut, date_fin, analyse,
        kpis=[('Ventes', 'nb_ventes'), ("Chiffre d'affaires (XOF)", 'ca_total'), ('Commandes', 'nb_commandes')])
    _set_column_widths(ws, 18)
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="rapport_distribution_{date_debut}_{date_fin}.xlsx"'
    wb.save(response)
    return response
